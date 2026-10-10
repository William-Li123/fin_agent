"""Explicit local document ingestion. Gold/annotation fields never become model input."""
from __future__ import annotations
import csv
import hashlib
from pathlib import Path
from finagent.io import read_json, save_json

FIELDS = {"id", "question", "pre_text", "post_text", "table"}


def ingest(path: str | Path, output_dir: str | Path, *, question: str = "") -> list[Path]:
    path, output_dir = Path(path).resolve(), Path(output_dir).resolve()
    if not path.is_file():
        raise ValueError("Input must be an existing local document")
    if path.stat().st_size > 25_000_000:
        raise ValueError("Document exceeds the 25 MB ingestion limit")
    content_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    base = "doc_" + content_hash[:20]
    suffix = path.suffix.lower()
    docs = []
    if suffix == ".json":
        doc = read_json(path)
        if not isinstance(doc, dict) or set(doc) != FIELDS:
            raise ValueError("JSON must contain exactly id,question,pre_text,post_text,table; use eval prepare for FinQA")
        docs.append(doc)
    elif suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        docs.append(dict(id=base, question=question, pre_text=[path.name], post_text=[], table=rows))
    elif suffix in {".txt", ".md"}:
        paragraphs = [p for p in path.read_text(encoding="utf-8-sig").split("\n\n") if p.strip()]
        docs.append(dict(id=base, question=question, pre_text=paragraphs, post_text=[], table=[]))
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF requires an unlocked local copy")
        paragraphs = [f"[PDF page {i+1}]\n{p.extract_text() or ''}" for i, p in enumerate(reader.pages)]
        if not any(len(p.split("\n", 1)[-1].strip()) for p in paragraphs):
            raise ValueError("PDF has no extractable text; OCR is not included")
        docs.append(dict(id=base, question=question, pre_text=paragraphs, post_text=[], table=[]))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for i, sheet in enumerate(workbook):
                table = [["" if c is None else str(c) for c in row] for row in sheet.iter_rows(values_only=True)]
                docs.append(dict(id=f"{base}_{i}", question=question, pre_text=[f"{path.name}: {sheet.title}"], post_text=[], table=table))
        finally:
            workbook.close()
    else:
        raise ValueError("Supported input formats: JSON, CSV, TXT, MD, PDF text, XLSX cached values")
    targets = []
    for index, doc in enumerate(docs):
        if question:
            doc["question"] = question
        target = output_dir / f"{base}_{index}.json"
        save_json(target, doc)
        targets.append(target)
    save_json(output_dir / "manifest.json", {"original_name": path.name, "original_sha256": content_hash,
        "parser": suffix, "documents": [str(p.name) for p in targets],
        "limitations": "PDF text extraction is not table OCR; XLSX formulas use cached values."})
    return targets

